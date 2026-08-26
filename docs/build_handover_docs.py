from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
SHOTS = DOCS / "screenshots"
USER_OUT = DOCS / "BNI_CHAPTER_PULSE_User_Manual.docx"
LOGIC_OUT = DOCS / "BNI_CHAPTER_PULSE_Logic_Handover.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
NAVY = RGBColor(11, 37, 69)
MUTED = RGBColor(85, 96, 116)
RED = RGBColor(224, 24, 40)
LIGHT = "E8EEF5"
PALE = "F4F6F9"
BORDER = "D7DBE2"
TODAY = "18 August 2026"


def _set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_border(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def _set_cell_width(cell, width_in):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_in * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def _set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(int(w * 1440) for w in widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            _set_cell_width(cell, widths[index])
            _set_cell_margins(cell)
            _set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _run_font(run, size=None, color=None, bold=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold


def _style_doc(doc, label):
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = sec.right_margin = sec.bottom_margin = sec.left_margin = Inches(1)
    sec.header_distance = Inches(0.492)
    sec.footer_distance = Inches(0.492)
    for name in ("Normal", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        st = doc.styles[name]
        st.font.size = Pt(size)
        st.font.color.rgb = color
        st.font.bold = True
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
    header = sec.header.paragraphs[0]
    header.text = label
    _run_font(header.runs[0], 9, MUTED)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.text = "BNI Chapter Pulse | Handover Pack"
    _run_font(footer.runs[0], 9, MUTED)


def _brand_title(doc, subtitle):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("BNI")
    _run_font(r, 24, RED, True)
    r2 = p.add_run(" CHAPTER PULSE")
    _run_font(r2, 24, NAVY, True)
    p2 = doc.add_paragraph(subtitle)
    _run_font(p2.runs[0], 12, MUTED)
    p2.paragraph_format.space_after = Pt(16)
    _kv_table(doc, [
        ("Prepared for", "Area Directors, Senior Directors, Chapter Directors, and project owner"),
        ("Project", "BNI CHAPTER PULSE"),
        ("Live site", "https://bni-pulse.netlify.app"),
        ("Prepared on", TODAY),
    ])
    doc.add_paragraph()


def _callout(doc, label, text, fill=PALE):
    table = doc.add_table(rows=1, cols=1)
    _table_geometry(table, [6.35])
    cell = table.cell(0, 0)
    _set_cell_shading(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(f"{label}: ")
    _run_font(r, 10.5, DARK_BLUE, True)
    _run_font(p.add_run(text), 10.5, NAVY)
    doc.add_paragraph()


def _kv_table(doc, rows, widths=(1.45, 4.9)):
    table = doc.add_table(rows=len(rows), cols=2)
    _table_geometry(table, list(widths))
    for idx, (key, value) in enumerate(rows):
        table.cell(idx, 0).text = key
        table.cell(idx, 1).text = value
        _set_cell_shading(table.cell(idx, 0), LIGHT)
        for cell in table.row_cells(idx):
            for para in cell.paragraphs:
                for run in para.runs:
                    _run_font(run, 10, NAVY)
    return table


def _matrix(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    _table_geometry(table, widths)
    for index, header in enumerate(headers):
        cell = table.cell(0, index)
        cell.text = header
        _set_cell_shading(cell, LIGHT)
        for run in cell.paragraphs[0].runs:
            _run_font(run, 9.3, NAVY, True)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = str(value)
        for index, cell in enumerate(cells):
            _set_cell_width(cell, widths[index])
            _set_cell_margins(cell)
            _set_cell_border(cell)
            for para in cell.paragraphs:
                for run in para.runs:
                    _run_font(run, 9.1, NAVY)
    doc.add_paragraph()


def _bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.188)
        p.paragraph_format.space_after = Pt(4)
        _run_font(p.add_run(item), 10.7, NAVY)


def _steps(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.188)
        p.paragraph_format.space_after = Pt(4)
        _run_font(p.add_run(item), 10.7, NAVY)


def _image(doc, filename, caption, width=6.25):
    path = SHOTS / filename
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(10)
    _run_font(cap.runs[0], 9, MUTED)


def build_user_manual():
    doc = Document()
    _style_doc(doc, "BNI Chapter Pulse - User Manual")
    _brand_title(doc, "User Manual with examples and dashboard screenshots")
    _callout(doc, "Purpose", "This manual explains how each user role works with the live BNI Chapter Pulse dashboard for weekly data, renewals, pipeline, planning, risk, and reports.")

    doc.add_heading("1. Login And Access", level=1)
    _image(doc, "01-login.png", "Login screen: choose your name, enter your private PIN, and sign in.")
    _steps(doc, [
        "Open https://bni-pulse.netlify.app.",
        "Select your name from the dropdown.",
        "Enter the PIN provided privately by the admin.",
        "Press Enter or click Login. If Session expired appears, refresh and sign in again.",
    ])
    _matrix(doc, ["Role", "Access", "Typical work"], [
        ["Area Director", "Whole region", "Regional targets, filters, renewals, retention, planning, reports, settings"],
        ["Senior Director", "Assigned chapters", "Manage assigned chapters, pipeline, renewals, risk, planning"],
        ["Chapter Director / DC", "Own chapters", "Weekly data, pipeline updates, renewal follow-up, chapter reports"],
    ], [1.45, 2.1, 2.8])

    doc.add_heading("2. Area Director Overview", level=1)
    _image(doc, "02-area-overview.png", "Area Director overview: targets, cumulative net added, filters, and chapter cards.")
    _bullets(doc, [
        "Total Members / Score shows current active members against the regional score target of 4404.",
        "Cumulative Net Added shows inductions minus drops for 30 days, 60 days, 90 days, 6 months, or 12 months.",
        "Support Scorecard Target is 696 per year, divided into 4 quarters of 174 inductions.",
        "Each quarter card shows inductions, drops, net movement, cumulative net, and remaining induction gap.",
        "Search Chapter / Director and Senior Director filter narrow the regional view.",
    ])
    _callout(doc, "Example", "If the selected period has 12 inductions and 5 drops, Cumulative Net Added is +7. Support scorecard progress counts the 12 inductions.")

    doc.add_heading("3. Weekly Data Entry", level=1)
    _bullets(doc, [
        "Select the chapter and meeting week before entering values.",
        "Enter opening members, inductions, drops, visitors, positive visitors, referrals, TYFCB, EOI Yes / Maybe / No, renewals, training, and success story where applicable.",
        "Members Close is calculated as Start Members + Inductions - Drops.",
        "Saving weekly data updates the DC, Senior Director, and Area Director dashboards because they read the same weeklyData records.",
    ])
    _callout(doc, "Example", "Ares starts with 32 members, inducts 2, and drops 1. Closing active members becomes 33 and weekly net add is +1.")

    doc.add_heading("4. Pipeline", level=1)
    _image(doc, "pipeline.png", "Pipeline module: track EOI Yes visitors through pipeline, inducted, or dropped status.")
    _bullets(doc, [
        "Use Pipeline for named visitor follow-up after EOI Yes.",
        "Marking a contact Inducted adds one induction into the matching weekly scorecard.",
        "The same contact cannot be double counted. Reopen removes the synced induction.",
        "Use chapter and stage filters to review pending, inducted, or dropped contacts.",
    ])

    doc.add_heading("5. Renewals And Retention", level=1)
    _image(doc, "renewals.png", "Renewals module: outstanding and upcoming renewals with 30, 60, and 90 day filters.")
    _bullets(doc, [
        "Renewals come from the uploaded Members Due report.",
        "The 30, 60, and 90 day buttons change the upcoming renewal window.",
        "Overdue renewals remain visible until marked complete.",
        "Mark Done writes a completion record and changes the row to Complete.",
    ])
    _image(doc, "retention.png", "Retention module: current month plus 3, 6, and 12 month analysis.")
    _bullets(doc, [
        "Retention is calculated from dues due versus renewals completed.",
        "Use this month, 3 months, 6 months, or 12 months to locate follow-up pressure.",
        "A chapter with 10 dues and 8 completed renewals shows 80% retention.",
    ])

    doc.add_heading("6. Planning And Risk Radar", level=1)
    _image(doc, "planning.png", "Planning module: targets, gaps, next month renewals, TLR, and health.")
    _bullets(doc, [
        "Planning compares current members with chapter targets and monthly goals.",
        "Next Month Renewals comes from Members Due records in the next calendar month that are not completed.",
        "Weekly target shows how many members per week are needed to close the gap.",
    ])
    _image(doc, "risk.png", "Risk Radar: chapters grouped by TLR and average weekly net add, with health score breakdown.")
    _bullets(doc, [
        "Thriving: TLR is 50 or above and average weekly net add is positive or zero.",
        "Structured but Shrinking: TLR is 50 or above, but average weekly net add is negative.",
        "Growing but Unstructured: TLR is below 50, but net add is positive or zero.",
        "Critical: TLR is below 50 and average weekly net add is negative.",
    ])

    doc.add_heading("7. Uploads, Reports, And Settings", level=1)
    _bullets(doc, [
        "TLR upload accepts PDF reports and merges chapter traffic light data.",
        "Members Due upload accepts Excel / XML spreadsheet files and powers Renewals and Retention.",
        "Miyagi upload accepts the wider Miyagi Excel report and powers member tracking views.",
        "Reports can be generated and exported as CSV for scorecard, renewal risk, funnel, Miyagi, retrospective, and operational reviews.",
    ])
    _image(doc, "settings.png", "Settings module: manage accounts, chapter assignments, PIN updates, team call day, and branding.")
    _bullets(doc, [
        "Admins can add or edit Area Directors, Senior Directors, Chapter Directors, SA roles, and viewers.",
        "Senior Directors and Chapter Directors must have the correct chapters assigned here.",
        "When a PIN changes, the confirmation message shows name, role, Senior DC relation, chapters, and new PIN for 10 seconds.",
        "After saving, the PIN field is cleared so the value is not left visible.",
    ])

    doc.add_heading("8. Quick Troubleshooting", level=1)
    _matrix(doc, ["Issue", "What to check"], [
        ["Incorrect PIN", "Confirm the selected name, then ask admin to reset the PIN."],
        ["Session expired", "Refresh the page. If needed, clear the site cache and sign in again."],
        ["Chapter missing", "Settings must include the chapter on that Senior Director or DC account."],
        ["Renewals empty", "Upload Members Due and confirm due dates are inside the selected window."],
        ["TLR or TYFCB missing", "Re-upload the correct TLR PDF and verify the parser layout."],
        ["Pipeline not reflected", "Confirm the contact is Inducted and the weekly scorecard row was created or updated."],
    ], [1.7, 4.65])
    _callout(doc, "Go-live habit", "Before sharing weekly numbers, review Overview, Renewals, Retention, Pipeline, Planning, Risk Radar, and Chapter Scorecard once with filters cleared.")
    doc.save(USER_OUT)


def build_logic_handover():
    doc = Document()
    _style_doc(doc, "BNI Chapter Pulse - Logic Handover")
    _brand_title(doc, "Logic and technical handover document")
    _callout(doc, "Purpose", "This document explains data storage, access control, dashboard calculations, parser behavior, and deployment checks for project handover.")

    doc.add_heading("1. System Architecture", level=1)
    _kv_table(doc, [
        ("Frontend", "Single-page dashboard in index.html, built and served by Netlify from dist."),
        ("Backend API", "Netlify Functions route /api/* through netlify/functions/api.js into src/api/handlers.js."),
        ("Database", "Firebase Firestore is the production source of truth."),
        ("Shared services", "src/services/firestore.js handles Firestore access; src/services/scope.js handles role scope."),
        ("Local preview", "Preview mode can use local/imported data without writing to live Firestore."),
    ])

    doc.add_heading("2. Firestore Data Model", level=1)
    _matrix(doc, ["Collection / document", "Purpose"], [
        ["members", "Leadership accounts, role, chapter assignment, reportsTo, chapterReportsTo, and bcrypt pinHash."],
        ["chapters", "Chapter master list: name, meeting day, target, Senior Director, Chapter Director, order."],
        ["weeklyData", "Weekly chapter scorecard records: members, inductions, drops, visitors, referrals, TYFCB, EOI, TLR."],
        ["visitorPipeline", "Named visitor pipeline contacts and stage: pipeline, inducted, or dropped."],
        ["renewalsDone", "Completion records after Mark Done."],
        ["attendance / miyagiMembers / activityLog", "Call attendance, Miyagi member tracking, and audit-style action logs."],
        ["meta/dues, meta/tlr, meta/monthlyTargets, meta/chapterGoals", "Uploaded report rows and planning targets."],
        ["meta/branding, meta/config", "Branding URLs, app config, and master PIN hash where used."],
    ], [2.15, 4.2])

    doc.add_heading("3. Authentication And Scoping", level=1)
    _bullets(doc, [
        "The login directory returns names and role metadata only; it never returns PINs or hashes.",
        "/api/auth/login validates the submitted PIN against the member bcrypt pinHash.",
        "Plaintext pin fields are rejected; login requires pinHash.",
        "Area Directors see the whole region. Viewers read the whole region but cannot write.",
        "Senior Directors, DC/CD, SA1, and SA2 are limited to chapters stored on their member record.",
        "Server-side scope filtering protects reads and writes so users cannot update chapters outside their access.",
    ])
    _callout(doc, "Credential rule", "Do not commit .env files, Firebase service account JSON, FIREBASE_SERVICE_ACCOUNT_BASE64, APP_JWT_SECRET, or member-pins.csv.")

    doc.add_heading("4. Chapter And Leadership Mapping", level=1)
    _bullets(doc, [
        "Only Snehal Patel and Yash Vasant should have Area Director whole-region access.",
        "All other region leaders are Senior Directors unless assigned a specific chapter/support role.",
        "Mapping is represented by members.chapters, members.reportsTo, members.chapterReportsTo, and chapters.seniorDirector / chapterDirector.",
        "Nachiket Patel and the removed team were intentionally removed from active region scope.",
        "CD and DC are treated as the same chapter-director level.",
    ])

    doc.add_heading("5. Weekly Scorecard Logic", level=1)
    _matrix(doc, ["Field", "Source / calculation"], [
        ["Members Start", "Entered by user, or seeded from latest prior close when pipeline creates a row."],
        ["Inductions", "Manual weekly entry plus pipeline contacts marked Inducted."],
        ["Drops", "Manual weekly entry."],
        ["Members Close", "max(0, Members Start + Inductions - Drops)."],
        ["Weekly Net Add", "Inductions - Drops for the selected week."],
        ["Monthly Net Add", "Last closing members in month minus first starting members in month."],
        ["TLR", "Most recent matching TLR parsed row for the chapter."],
        ["Reports visibility", "All dashboards read weeklyData, then filter by user scope."],
    ], [1.75, 4.6])
    _callout(doc, "Example", "Start 32 + 2 inductions - 1 drop = Members Close 33 and Weekly Net Add +1.")

    doc.add_heading("6. Pipeline Sync Logic", level=1)
    _bullets(doc, [
        "Each contact has id, chapter, stage, stageChangedDate, and stageChangedBy.",
        "Changing stage to inducted increments weeklyData.inductions by one for the matching week.",
        "The id is stored in weeklyData.pipelineInductedIds to prevent double counting.",
        "Reopening an inducted contact removes that id and decrements the induction.",
        "Members Close is recalculated immediately after the induction adjustment.",
    ])

    doc.add_heading("7. Area Director Targets", level=1)
    _matrix(doc, ["Metric", "Logic"], [
        ["Total Members / Score", "Sum latest closing members for visible chapters, compared with 4404."],
        ["Support Scorecard Target", "696 inductions per support year."],
        ["Quarter Target", "696 / 4 = 174 inductions per quarter."],
        ["Quarter Period", "Support year starts from current month; each quarter covers 3 calendar months."],
        ["Quarter Actual", "Total inductions in that quarter for visible chapters."],
        ["Quarter Drops / Net", "Drops in quarter; net = inductions - drops."],
        ["Cumulative Net Added", "Inductions minus drops over 30d, 60d, 90d, 6m, or 12m."],
    ], [2.1, 4.25])

    doc.add_heading("8. Renewals And Retention", level=1)
    _bullets(doc, [
        "Members Due data is stored in meta/dues.members.",
        "Renewals shows outstanding overdue records plus upcoming records in the selected 30, 60, or 90 day window.",
        "Mark Done writes renewalsDone and changes the row to Complete.",
        "Retention calculates completed renewals divided by dues due for the selected current month, 3 month, 6 month, or 12 month period.",
    ])

    doc.add_heading("9. Risk Radar And Health Score", level=1)
    _matrix(doc, ["Component", "Points", "Rule"], [
        ["Compliance", "15", "Last 4 weeks submission count / 4 * 15, rounded."],
        ["Net Add", "20", "Pass/fail: last 3 calendar months each need positive monthly net or at least one induction."],
        ["TLR", "20", "Pass/fail: TLR >= 50 earns 20."],
        ["Conversion", "20", "Average recent conversion: >=50 earns 20, >=30 earns 14, >=15 earns 8, else 0."],
        ["EOI Follow-through", "25", "Pending EOI backlog: 0 earns 25, <=2 earns 18, <=5 earns 10, else 0."],
        ["Grade", "A/B/C/D", "A >=80, B >=60, C >=40, D below 40."],
    ], [1.55, 0.7, 4.1])
    _bullets(doc, [
        "Risk quadrant uses TLR threshold 50 and average weekly net add.",
        "Thriving means high TLR and non-negative net add.",
        "Critical means low TLR and negative net add.",
    ])

    doc.add_heading("10. Parser And Upload Logic", level=1)
    _matrix(doc, ["Upload", "Expected file", "Important fields"], [
        ["TLR", "PDF chapter traffic light reports", "Score, size, growth, retention, referrals, visitors, conversion, absenteeism, TYFCB where present."],
        ["Members Due", "Excel or XML spreadsheet report", "Chapter, member name, industry, type, status, due / renewal date."],
        ["Miyagi", "Excel sheet with wider member columns", "Chapter and member tracking columns used by Miyagi / 1YRC reports."],
    ], [1.25, 1.9, 3.2])
    _callout(doc, "Parser maintenance", "If PDF layout remains stable, TLR parsing should remain stable. If an Excel format changes, test locally and update header matching before pushing live.")

    doc.add_heading("11. Deployment And Operations", level=1)
    _bullets(doc, [
        "Run npm run check and npm run build before deployment.",
        "Netlify build command is npm run build and publish directory is dist.",
        "Production variables include Firebase project values, FIREBASE_SERVICE_ACCOUNT_BASE64, APP_JWT_SECRET, and NODE_ENV=production.",
        "Use Netlify Functions for backend routes instead of exposing Firebase Admin credentials in the browser.",
        "After deploy, test login, data entry save, pipeline inducted sync, renewal Mark Done, AD filters, retention filters, planning, risk, uploads, and one CSV export.",
    ])
    _matrix(doc, ["Smoke test", "Expected result"], [
        ["Login as Area Director", "Whole region visible; Senior Director and chapter filters available."],
        ["Login as Senior Director", "Only assigned chapters appear in dropdowns and reports."],
        ["Login as DC/CD", "Only assigned chapter or chapters appear."],
        ["Change PIN", "10 second confirmation shows name, role, Senior DC relation, chapters, and new PIN; field clears."],
        ["Mark pipeline Inducted", "Pipeline stage changes and weekly scorecard induction/member close update."],
        ["Mark renewal Done", "Renewal row becomes Complete and retention numbers update."],
    ], [2.0, 4.35])

    doc.add_heading("12. Handover Checklist", level=1)
    _bullets(doc, [
        "Confirm final live URL and Netlify site owner.",
        "Confirm Firebase project owner and backup admin access.",
        "Store production environment variables in Netlify only, not in Git.",
        "Keep credential handover separate from the public user manual.",
        "Export a backup of Firestore collections before large imports.",
        "Keep sample TLR, Members Due, and Miyagi files for parser regression testing.",
    ])
    doc.save(LOGIC_OUT)


if __name__ == "__main__":
    build_user_manual()
    build_logic_handover()
    print(USER_OUT)
    print(LOGIC_OUT)
